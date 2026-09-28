"""
Parsers for the scenario/topic taxonomy files -- `List of Scenarios - Tanisha.docx` (the largest
single file in the whole convo-data drop) and `Anxiety Situations.docx` (a simple reference table).
See documnets/understanding/21_NEW_CONTENT_INGESTION_PLAN.md for the full file-by-file reasoning.

CONFIRMED SENSITIVE CONTENT: `List of Scenarios - Tanisha.docx` has a "Fetishes" top-level category
covering all 8 DSM-5-TR paraphilic disorders (including Pedophilic and Sexual Sadism Disorder,
framed as non-offending help-seeking content). Per
documnets/understanding/15_OPEN_QUESTIONS_AND_BLOCKERS.md #26, this needs an explicit clinical +
legal governance decision before going live. This parser does NOT drop that content (dropping it
silently would be worse than flagging it) -- every row from that category is tagged
`pending_governance_review=True`, and the loader/CLI must exclude rows with that flag from
anything that becomes a live, queryable index until the decision is made.
"""

from pathlib import Path

from docx import Document

from app.ingestion.parsers.docx_common import get_table_data

_GOVERNANCE_HOLD_CATEGORIES = {"Fetishes"}


def parse_list_of_scenarios(path: str | Path) -> list[dict]:
    """
    Confirmed real structure: 9 `Title`-style top-level categories (the 9th, "Scenarios That Were
    Left", is itself a container for 10 further `Heading 1` sub-categories -- an unmerged second
    authoring pass, per documnets/understanding/18_NEW_CONVO_DATA_CATALOG.md §4b).

    Within a category, each topic starts with a BOLD "Normal"-style paragraph (confirmed via
    `paragraph.runs[0].bold`), optionally followed by a "Scenario: ..." one-line clinical gloss and
    a literal "Example Statements" marker (confirmed present in the "Relationship Issues" category,
    confirmed ABSENT in "Particular Mental Health Issues" and "Scenarios That Were Left" -- the
    document is not internally consistent here). All non-bold lines after a topic's bold header,
    until the next bold line, are treated as phrases.

    KNOWN SIMPLIFICATION, documented rather than silently applied: some sections (e.g. "Particular
    Mental Health Issues") have a THIRD level of nesting -- short, non-bold sub-labels like
    "Procrastination" or "Difficulty focusing" sitting between the bold topic and its phrases. This
    parser does not attempt to detect that third level; those sub-labels are stored as ordinary
    phrases alongside the real ones. This is a deliberate simplification (detecting them reliably
    would need a fragile heuristic), not an oversight -- see
    documnets/understanding/21_NEW_CONTENT_INGESTION_PLAN.md for the reasoning.
    """
    document = Document(str(path))
    records: list[dict] = []

    current_category: str | None = None
    current_topic: str | None = None
    current_scenario_desc: str | None = None

    for p in document.paragraphs:
        text = p.text.strip()
        style = p.style.name if p.style else ""
        is_bold = bool(p.runs and p.runs[0].bold)

        if style in ("Title", "Heading 1"):
            current_category = text
            current_topic = None
            current_scenario_desc = None
            continue

        if not text or current_category is None:
            continue

        if is_bold:
            current_topic = text
            current_scenario_desc = None
            continue

        if text.lower().startswith("scenario:"):
            current_scenario_desc = text[len("scenario:") :].strip()
            continue

        if text == "Example Statements":
            continue  # literal section marker, not content

        if current_topic is None:
            continue  # stray text before the first topic in this category

        records.append(
            {
                "category": current_category,
                "topic": current_topic,
                "scenario_description": current_scenario_desc,
                "phrase": text,
                "pending_governance_review": current_category in _GOVERNANCE_HOLD_CATEGORIES,
                "source_file": Path(path).name,
            }
        )
    return records


def parse_anxiety_situations(path: str | Path) -> list[dict]:
    """
    Confirmed real structure: a single 3-column table (`Sr. No | Broad Situation | Subtopics`),
    36 rows including header. This is a pure taxonomy/index, not phrase content -- see
    documnets/understanding/19_DATA_USAGE_MAP.md Bucket 2a entry for `Anxiety Situations.docx`.
    """
    document = Document(str(path))
    if not document.tables:
        return []

    rows = get_table_data(document.tables[0])
    records = []
    for row in rows[1:]:  # skip header row
        if len(row) < 3 or not row[1].strip():
            continue
        records.append(
            {
                "broad_situation": row[1].strip(),
                "subtopics_raw": row[2].strip(),
                "source_file": Path(path).name,
            }
        )
    return records
