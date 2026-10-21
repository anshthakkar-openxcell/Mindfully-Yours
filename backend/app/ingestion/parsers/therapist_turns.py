"""
Parses the therapists' full "User: / AI:" dialogue files into (user message -> therapist reply)
pairs for kb_therapist_turns -- the data behind the therapist-first response approach
(documnets/understanding/32_THERAPIST_FIRST_RESPONSE_APPROACH.md).

Every format fact below was confirmed by scanning the real decrypted files (22 of the 46 use this
dialogue shape; no two are identical):
- Speaker prefixes vary: User / Client / Parent on the user side, AI / Ai / Al / "AI." / A on the AI
  side (all real, several are typos in the source).
- Conversation boundaries are NOT uniform: Heading 1/2/4 styles in some files, "Conversation 3 (Tier 1)
  - title" lines in others, numbered lines ("1a: ...") in others, nothing at all in some. So the rule
  is generic: any non-dialogue line between two dialogue runs ends the conversation and becomes
  header text for the next one. This can split a conversation in two (safe: we lose a pair, never
  invent one) but never merges two different conversations.
- A bullet/list line directly after an AI turn is that turn's continuation (a list of tips), not a
  boundary.
- `*...*` lines are drafting notes ("*Detailed Grounding methods*") standing in for text the
  therapist hasn't written yet -- the reply is held, never served with the note missing.

SAFETY POSTURE (this is what makes the approach MORE accurate, not just more therapist-flavoured --
doc 32 Gate 2): when in doubt a row is stored but marked match_eligible=False with a hold_reason.
Never served: placeholders/drafting notes, Tier-3/crisis content (the red-flag interlock owns that
path, verbatim), files on governance hold (15_OPEN_QUESTIONS_AND_BLOCKERS.md #26), and very short
user messages ("yes", "thanks") that would match anything.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path

from docx import Document

USER_SPEAKER_RE = re.compile(r"^\s*(?:User|Client|Parent|Female)\s*[:\-–]\s*(?P<text>.*)$", re.IGNORECASE | re.DOTALL)
AI_SPEAKER_RE = re.compile(r"^\s*(?:AI Companion|AI|Al)\s*(?:[:\-–]|\.(?=\s))\s*(?P<text>.*)$", re.IGNORECASE | re.DOTALL)
# Bare "A:" is a real (rare) typo in the source -- colon only, a bare "A -" / "A." is too ambiguous.
AI_BARE_A_RE = re.compile(r"^\s*A\s*:\s*(?P<text>.*)$", re.DOTALL)

TIER_RE = re.compile(r"Tier\s*:?\s*([123])", re.IGNORECASE)
TIER_ONLY_RE = re.compile(r"^\s*Tier\s*[123]\b.{0,30}$", re.IGNORECASE)  # "Tier 3", "Tier 3 (High Concern)"
# A numbered line like "38. User is upset due to her grandparents" is the NEXT scenario's heading,
# not a tip in the previous reply (confirmed in Miscellaneous Conversations1 and the Anam files --
# before this guard, ~100 replies had a stray scenario heading glued onto the end, which would have
# been spoken to a real user word for word).
SCENARIO_HEADING_RE = re.compile(r"^\s*\d+[.)]\s+(?:the\s+|a\s+|an\s+)?(?:user|client|parent|caregiver)", re.IGNORECASE)
BULLET_RE = re.compile(r"^\s*(?:[•\-–]|\d+[.)]|[a-z][.)])\s+\S")
DRAFTING_NOTE_RE = re.compile(r"^\s*\*[^*]+\*\s*$")
PLACEHOLDER_RE = re.compile(
    r"\[[^\]]+\]"  # [technique], [xyz issue]
    r"|\((?:name|insert|add|could be|to be|specify|city|hometown)[^)]*\)"
    r"|\bxyz\b",
    re.IGNORECASE,
)
PROTOCOL_MARKER_RE = re.compile(r"^\s*(?:initiate|follow)\b.*protocols?\s*$|prescribed protocol", re.IGNORECASE)
CONVERSATION_PREFIX_RE = re.compile(r"^\s*Conversations?\s*\d*[a-z]?\s*(?:\([^)]*\))?\s*[:\-–.]?\s*", re.IGNORECASE)
CRISIS_TITLE_RE = re.compile(
    r"suicid|self[- ]?harm|\bkill|homicid|harm to (?:self|someone)|crisis|rape|molest|abus|overdose",
    re.IGNORECASE,
)

MIN_USER_WORDS = 4  # below this ("yes", "thank you", "I don't know") a vector match is meaningless
PRIOR_CONTEXT_EXCHANGES = 2

# Never ingested: Career Conversation Sets 1.docx is 1 template repeated 500x -- already
# deduplicated into its real parts (see 19_DATA_USAGE_MAP.md), ingesting the raw file would add 500
# near-identical rows that all "match" every career message.
SKIP_FILES = {"Career Conversation Sets 1.docx"}
# Stored but never eligible -- see 15_OPEN_QUESTIONS_AND_BLOCKERS.md #26.
HELD_FILES = {"22-8_Offline Work Week 2.docx": "governance_hold_paraphilic_content_#26"}

# The client's curated "best of" file -- rows from it get is_preferred=True (a match-time tie-breaker).
PREFERRED_FILES = {"Best Conversations- Final.docx"}

KNOWN_AUTHORS = ("Anam", "Ankita", "Shreya", "Tanisha")


def _classify(text: str) -> tuple[str | None, str]:
    """('user' | 'ai' | None, text-without-prefix)."""
    if m := USER_SPEAKER_RE.match(text):
        return "user", m.group("text").strip()
    if m := AI_SPEAKER_RE.match(text):
        return "ai", m.group("text").strip()
    if m := AI_BARE_A_RE.match(text):
        return "ai", m.group("text").strip()
    return None, text


_NUMBERED_RE = re.compile(r"^\s*(\d+)[.)]\s+")


def _is_numbered_heading(line: str, reply_so_far: str, next_is_user: bool) -> bool:
    """
    A numbered line sitting right before a User turn, that does not extend a 1,2,3 sequence already
    inside this reply, is the next scenario's heading ("13. Emotional dys-regulation...") -- not a
    tip. A genuine numbered tip list is either followed by more list lines, or continues a sequence.
    """
    m = _NUMBERED_RE.match(line)
    if not m or not next_is_user:
        return False
    last_line = reply_so_far.rsplit("\n", 1)[-1]
    prev = _NUMBERED_RE.match(last_line)
    return not (prev and int(prev.group(1)) + 1 == int(m.group(1)))


def _has_stray_numbering(reply: str) -> bool:
    """True if the reply has numbered lines that are not a clean 1,2,3... run -- the signature of a
    scenario heading ("7. Losing Something Important") that leaked onto the end of the reply. Catch-all
    behind the structural checks above: a reply we cannot trust is held, never served."""
    numbers = [int(m.group(1)) for line in reply.split("\n") if (m := _NUMBERED_RE.match(line))]
    return bool(numbers) and numbers != list(range(1, len(numbers) + 1))


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


def _author_for(path: Path) -> str | None:
    haystack = " ".join(path.parts[-4:])
    for name in KNOWN_AUTHORS:
        if name.lower() in haystack.lower():
            return name
    return None


def _clean_title(line: str) -> str:
    return CONVERSATION_PREFIX_RE.sub("", line).strip(" -–:.") or line.strip()


def parse_therapist_dialogue_file(path: Path, *, root: Path) -> list[dict]:
    """One .docx -> list of kb_therapist_turns row dicts (no embedding yet)."""
    document = Document(str(path))
    source_file = str(path.relative_to(root))
    author = _author_for(path)
    file_hold = HELD_FILES.get(path.name)

    # --- Pass 1: group paragraphs into conversations of ordered (speaker, text) turns ----------
    conversations: list[dict] = []
    current: dict | None = None
    pending_headers: list[str] = []
    current_section: str | None = None

    def close_current() -> None:
        nonlocal current
        if current and current["turns"]:
            conversations.append(current)
        current = None

    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        style = paragraph.style.name or ""
        next_is_user = idx + 1 < len(paragraphs) and _classify(paragraphs[idx + 1].text.strip())[0] == "user"
        speaker, body = _classify(text)

        if speaker is not None:
            if current is None:
                current = {
                    "headers": pending_headers,
                    "section": current_section,
                    "turns": [],
                    "markers": [],
                }
                pending_headers = []
            if current["turns"] and current["turns"][-1]["speaker"] == speaker:
                # Consecutive lines by the same speaker are one turn (a long reply split by paragraphs).
                current["turns"][-1]["text"] += "\n" + body
            else:
                current["turns"].append({"speaker": speaker, "text": body, "drafting_note": False})
            continue

        # Non-dialogue line.
        if (
            current
            and current["turns"]
            and current["turns"][-1]["speaker"] == "ai"
            and (style.startswith("List") or BULLET_RE.match(text))
            and not SCENARIO_HEADING_RE.match(text)
            and not _is_numbered_heading(text, current["turns"][-1]["text"], next_is_user)
        ):
            current["turns"][-1]["text"] += "\n" + text  # tip list continuing the AI's reply
            continue
        if current and current["turns"] and current["turns"][-1]["speaker"] == "ai" and DRAFTING_NOTE_RE.match(text):
            current["turns"][-1]["drafting_note"] = True  # reply is incomplete in the source
            continue

        if PROTOCOL_MARKER_RE.search(text):
            # "Initiate crisis protocol" / "Initiate matching protocol" is the routing OUTCOME written
            # AFTER the conversation it describes (confirmed in the ED and Addictions files, where it
            # follows the dialogue and is followed by the next topic) -- so it belongs to the
            # conversation just finished, never to the next one. Getting this backwards flagged
            # unrelated topics (Weed, Gaming) as crisis.
            target = current if current else (conversations[-1] if conversations else None)
            if target is not None:
                target["markers"].append(text)
            close_current()
            continue

        close_current()
        if style.startswith("Heading") or style == "Title":
            current_section = text
        pending_headers.append(text)
    close_current()

    # --- Pass 2: turn each conversation into (user -> AI) pair rows ---------------------------
    rows: list[dict] = []
    slug = re.sub(r"[^a-z0-9]+", "_", path.stem.lower()).strip("_")
    for conv_no, conv in enumerate(conversations, start=1):
        headers = conv["headers"]
        tier = next((int(m.group(1)) for h in reversed(headers) if (m := TIER_RE.search(h))), None)
        marker_lines = conv["markers"]
        title_candidates = [h for h in headers if not TIER_ONLY_RE.match(h)]
        scenario = _clean_title(title_candidates[-1]) if title_candidates else None
        crisis = (
            tier == 3
            or any("crisis" in m.lower() for m in marker_lines)
            or bool(scenario and CRISIS_TITLE_RE.search(scenario))
            or bool(conv["section"] and CRISIS_TITLE_RE.search(conv["section"]))
        )
        conv_id = f"{slug}:{conv_no}"
        turns = conv["turns"]

        pair_idx = 0
        for i, turn in enumerate(turns):
            if turn["speaker"] != "ai" or i == 0 or turns[i - 1]["speaker"] != "user":
                continue
            user_text = turns[i - 1]["text"].strip()
            reply_text = turn["text"].strip()
            if not user_text or not reply_text:
                continue

            prior = turns[max(0, i - 1 - 2 * PRIOR_CONTEXT_EXCHANGES) : i - 1]
            prior_context = "\n".join(f"{'User' if t['speaker'] == 'user' else 'AI'}: {t['text']}" for t in prior) or None

            hold_reason = file_hold
            if hold_reason is None and turn["drafting_note"]:
                hold_reason = "reply_has_drafting_note"
            if hold_reason is None and PLACEHOLDER_RE.search(reply_text):
                hold_reason = "reply_has_placeholder"
            if hold_reason is None and _has_stray_numbering(reply_text):
                hold_reason = "reply_may_contain_stray_heading"
            if hold_reason is None and crisis:
                hold_reason = "crisis_or_tier3_handled_by_interlock"
            if hold_reason is None and len(user_text.split()) < MIN_USER_WORDS:
                hold_reason = "user_message_too_short_to_match"

            content_hash = hashlib.sha256(f"{_norm(user_text)}||{_norm(reply_text)}".encode()).hexdigest()
            rows.append(
                {
                    "conversation_id": conv_id,
                    "turn_index": pair_idx,
                    "user_text": user_text,
                    "prior_context": prior_context,
                    "reply_text": reply_text,
                    "scenario": scenario,
                    "section": conv["section"],
                    "tier": tier,
                    "emotion": None,
                    "layer": None,
                    "author": author,
                    "source_file": source_file,
                    "is_opening": i == 1,
                    "match_eligible": hold_reason is None,
                    "hold_reason": hold_reason,
                    "approved": False,
                    "is_preferred": path.name in PREFERRED_FILES,
                    "content_hash": content_hash,
                    "language": "en",
                }
            )
            pair_idx += 1
    return rows


def parse_therapist_turns(convo_dir: Path) -> tuple[list[dict], dict]:
    """
    Parse every dialogue .docx under `convo_dir` ('MYPL - Mindfully Yours Work/'). Returns
    (deduplicated rows, stats). Files with no User/AI lines (taxonomies, question banks) simply yield
    zero rows -- they're handled by their own parsers.
    """
    rows: list[dict] = []
    per_file: Counter[str] = Counter()
    for path in sorted(convo_dir.rglob("*.docx")):
        if path.name in SKIP_FILES or path.name.startswith("~$"):
            continue
        file_rows = parse_therapist_dialogue_file(path, root=convo_dir)
        per_file[str(path.relative_to(convo_dir))] = len(file_rows)
        rows.extend(file_rows)

    seen: set[str] = set()
    unique: list[dict] = []
    for row in rows:
        if row["content_hash"] in seen:
            continue
        seen.add(row["content_hash"])
        unique.append(row)

    stats = {
        "pairs_found": len(rows),
        "duplicates_dropped": len(rows) - len(unique),
        "match_eligible": sum(r["match_eligible"] for r in unique),
        "held": Counter(r["hold_reason"] for r in unique if r["hold_reason"]),
        "per_file": per_file,
    }
    return unique, stats
