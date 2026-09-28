"""
Parser for `Parenting_Child Related.docx`'s Guardrail table -- a ready-made concern-category to
referral-urgency mapping for child-safety concerns. See
documnets/understanding/19_DATA_USAGE_MAP.md Bucket 3 ("Parenting_Child Related.docx -- mainly
valuable for its Guardrail table, a ready-made escalation-rule structure").

This is Bucket 3 content -- a deterministic lookup, never embedded or searched. Confirmed real
structure: the document has 4 tables total; the Guardrail table is the 4th (index 3), a simple
2-column `Category | Refer?` table, 19 rows including header. `Refer?` values confirmed to include
at least `Yes`, `Immediate`, and `Immediate/emergency` -- three distinct urgency levels, not a
simple yes/no.
"""

from pathlib import Path

from docx import Document

from app.ingestion.parsers.docx_common import get_table_data

GUARDRAIL_TABLE_INDEX = 3  # confirmed via direct inspection -- the 4th table in the document


def parse_guardrail_table(path: str | Path) -> list[dict]:
    document = Document(str(path))
    if len(document.tables) <= GUARDRAIL_TABLE_INDEX:
        raise ValueError(
            f"Expected at least {GUARDRAIL_TABLE_INDEX + 1} tables in {path}, "
            f"found {len(document.tables)} -- the document structure may have changed."
        )

    rows = get_table_data(document.tables[GUARDRAIL_TABLE_INDEX])
    records = []
    for row in rows[1:]:  # skip header row
        if len(row) < 2 or not row[0].strip():
            continue
        records.append(
            {
                "concern_category": row[0].strip(),
                "referral_urgency": row[1].strip(),  # e.g. 'Yes', 'Immediate', 'Immediate/emergency'
                "source_file": Path(path).name,
            }
        )
    return records
