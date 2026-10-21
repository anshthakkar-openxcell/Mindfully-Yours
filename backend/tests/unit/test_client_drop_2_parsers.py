"""Parsers for client drop #2 -- welcome messages and the Situational Routing sheet."""

import openpyxl
from docx import Document

from app.ingestion.parsers import parse_situational_routing, parse_welcome_responses


def test_welcome_responses_strip_numbering_dedupe_and_flag_non_greeting(tmp_path):
    d = Document()
    for line in [
        "AI WELCOME RESPONSES",
        "Hi user, thank you for sharing this with me.",
        "5. Hi user, it's nice to have you here. How can I help?",
        "Hi user, it's nice to have you here. How can I help?",
    ]:
        d.add_paragraph(line)
    path = tmp_path / "AI WELCOME RESPONSES.docx"
    d.save(str(path))
    rows = parse_welcome_responses(path)
    assert [r["text"] for r in rows] == [
        "Hi user, thank you for sharing this with me.",
        "Hi user, it's nice to have you here. How can I help?",
    ]
    assert [r["approved"] for r in rows] == [False, True]  # the first is not a greeting


def test_situational_routing_groups_and_missing_strategy(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Situation ID", "Situations", "T1", "T2", "T3", "Self-Care"])
    ws.append(["Work / career / study"])
    ws.append(["SID -001", "Losing a job", "t1", "t2", "t3", "Give yourself time."])
    ws.append(["SID -002", "Fear of being fired", "t1", "t2", "t3", None])
    ws.append(["Money / financial"])
    ws.append(["SID -010", "Debt", "t1", "t2", "t3", "Make a plan."])
    path = tmp_path / "Situational Routing.xlsx"
    wb.save(path)
    rows = parse_situational_routing(path)
    assert [(r["situation_id"], r["category"]) for r in rows] == [
        ("SID-001", "Work / career / study"),
        ("SID-002", "Work / career / study"),
        ("SID-010", "Money / financial"),
    ]
    assert rows[1]["self_care_strategy"] is None  # a missing strategy stays NULL, never invented
