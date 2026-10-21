"""
Parser tests for app.ingestion.parsers.therapist_turns -- built on tiny synthetic .docx files that
reproduce the real quirks confirmed in the therapists' files (see that module's docstring), each one
a bug that actually happened while ingesting the real data.
"""

from pathlib import Path

from docx import Document

from app.ingestion.parsers.therapist_turns import parse_therapist_dialogue_file


def _doc(tmp_path: Path, lines: list[tuple[str, str | None]], name: str = "Sample.docx") -> Path:
    d = Document()
    for text, style in lines:
        d.add_paragraph(text, style=style) if style else d.add_paragraph(text)
    path = tmp_path / name
    d.save(str(path))
    return path


def _parse(tmp_path: Path, lines, name: str = "Sample.docx") -> list[dict]:
    return parse_therapist_dialogue_file(_doc(tmp_path, lines, name), root=tmp_path)


def test_basic_pairs_and_speaker_variants(tmp_path):
    rows = _parse(
        tmp_path,
        [
            ("Conversation 1 (Tier 1) - Feeling homesick", None),
            ("Client: I moved here last month and I feel really lonely.", None),
            ("Ai: That sounds hard. What has been the toughest part?", None),
            ("Client: Not knowing anyone at all in this city.", None),
            ("AI. Would you like some ideas to meet people?", None),
        ],
    )
    assert [r["reply_text"] for r in rows] == [
        "That sounds hard. What has been the toughest part?",
        "Would you like some ideas to meet people?",
    ]
    assert rows[0]["is_opening"] and not rows[1]["is_opening"]
    assert rows[0]["tier"] == 1 and rows[0]["scenario"] == "Feeling homesick"
    assert all(r["approved"] is False for r in rows)  # nothing is servable until a therapist approves


def test_protocol_marker_belongs_to_the_conversation_before_it(tmp_path):
    rows = _parse(
        tmp_path,
        [
            ("Weed", None),
            ("User: I have started using weed a few times a week to relax.", None),
            ("AI: Thanks for telling me. How is it affecting your day?", None),
            ("Initiate crisis protocol", None),
            ("Gaming", None),
            ("User: I play games all night and my sleep is ruined lately.", None),
            ("AI: That sounds exhausting. How long has it been going on?", None),
        ],
    )
    weed, gaming = rows
    assert weed["hold_reason"] == "crisis_or_tier3_handled_by_interlock"
    # The bug this guards: the marker used to be attached to the NEXT topic, wrongly holding it.
    assert gaming["hold_reason"] is None and gaming["match_eligible"]


def test_scenario_heading_never_leaks_into_reply(tmp_path):
    rows = _parse(
        tmp_path,
        [
            ("User: I feel anxious about my exams and cannot focus at all.", None),
            ("AI: Thank you for sharing. How long has this been going on?", None),
            ("User: A few weeks now, and it is getting worse each day.", None),
            ("AI: That is good to know.", None),
            ("7. Losing Something Important", None),
            ("User: I lost my wallet today and I am so upset about it.", None),
            ("AI: I am sorry to hear that. What was in it?", None),
        ],
    )
    assert "7." not in rows[1]["reply_text"] and "Losing" not in rows[1]["reply_text"]


def test_drafting_note_and_placeholder_are_held(tmp_path):
    rows = _parse(
        tmp_path,
        [
            ("User: I keep spiralling and cannot breathe properly at night.", None),
            ("AI: Let's try a grounding exercise together.", None),
            ("*Detailed grounding method*", None),
            ("User: Okay, I can try that tonight before bed.", None),
            ("AI: Try the [technique] we talked about before sleeping.", None),
        ],
    )
    assert rows[0]["hold_reason"] == "reply_has_drafting_note"
    assert rows[1]["hold_reason"] == "reply_has_placeholder"
    assert not any(r["match_eligible"] for r in rows)


def test_short_user_message_and_governance_file_are_held(tmp_path):
    rows = _parse(
        tmp_path,
        [("User: Yes please.", None), ("AI: Great, here is the first step for you today.", None)],
    )
    assert rows[0]["hold_reason"] == "user_message_too_short_to_match"

    held = _parse(
        tmp_path,
        [("User: I have been struggling with this for a long time.", None), ("AI: Tell me more about it.", None)],
        name="22-8_Offline Work Week 2.docx",
    )
    assert held[0]["hold_reason"].startswith("governance_hold")


def test_female_and_ai_companion_speaker_labels(tmp_path):
    rows = _parse(
        tmp_path,
        [
            ("AI Companion: Hey, I'm your AI companion. What would you like to talk about?", None),
            ("User: I thought I would tell you how my day went, it was strange.", None),
            ("AI Companion: Strange days can be interesting. What happened?", None),
            ("Female: My parents keep pressuring me to get married this year.", None),
            ("AI: That sounds stressful. How are you handling it?", None),
        ],
    )
    assert [r["reply_text"] for r in rows] == [
        "Strange days can be interesting. What happened?",
        "That sounds stressful. How are you handling it?",
    ]


def test_best_conversations_file_is_preferred(tmp_path):
    rows = _parse(
        tmp_path,
        [("User: I feel sad most of the time and I do not know why.", None), ("AI: How long has this been going on?", None)],
        name="Best Conversations- Final.docx",
    )
    assert rows[0]["is_preferred"] is True
